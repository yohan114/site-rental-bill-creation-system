#!/usr/bin/env python3
"""Export the rates embedded in index.html to an Excel rate card:
data/Fleet_Rental_Rate_Matrix_2026.xlsx

The workbook mirrors exactly what the invoice system charges — all three
rate-basis tiers (FULLY WET / WET driver-only / DRY) across Hourly, Daily and
Per-KM billing for every fleet unit, plus the portable equipment day rates.

Run after every `build_fleet_data.py --write` to keep the Excel in sync:
    python3 tools/export_rate_matrix.py

Requires: pip install openpyxl
"""

import argparse
import json
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK = "14110F"
ORANGE = "F58220"
ORANGE_SOFT = "FEF3E7"
GREEN = "2E9E4F"
GREEN_SOFT = "E8F5EC"
GRAY = "8A8278"
GRAY_SOFT = "F0EBE3"
LINE = "E4DFD7"

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_data(html_path):
    html = open(html_path, encoding="utf-8").read()
    fleet = json.loads(re.search(r"const FLEET_MACHINES = (\[.*?\]);", html, re.S)
                       .group(1).replace("<\\/", "</"))
    portable = json.loads(re.search(r"const PORTABLE_EQUIPMENT = (\[.*?\]);", html, re.S)
                          .group(1).replace("<\\/", "</"))
    return fleet, portable


def style_header(ws, row, cols, fills):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Calibri", bold=True, size=10,
                         color="FFFFFF" if fills.get(c, INK) in (INK, ORANGE, GREEN, GRAY) else INK)
        cell.fill = PatternFill("solid", fgColor=fills.get(c, INK))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def sheet_readme(wb):
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    lines = [
        ("2026 FLEET RATE MATRIX — Edward & Christie", 14, True, ORANGE),
        (f"Generated {date.today():%d-%b-%Y} from the rates embedded in index.html "
         "(the invoice system and this workbook always quote the same prices).", 10, False, None),
        ("", 10, False, None),
        ("RATE BASIS TIERS", 11, True, None),
        ("  FULLY WET — vehicle + driver + fuel (the Fleet Pricing sheet 'WET' rates)", 10, False, None),
        ("  WET — vehicle + driver, customer supplies fuel "
         "(derived: DRY + operator cost/h × 1.2, rounded to Rs. 50, clamped into [DRY, FULLY WET])", 10, False, None),
        ("  DRY — vehicle only (the Fleet Pricing sheet 'DRY' rates)", 10, False, None),
        ("", 10, False, None),
        ("BILLING METHODS", 11, True, None),
        ("  Hourly (min 120 hrs) · Daily = 4-hour day block (min 1 day) · "
         "Per-KM (min 3,000 km, transport units only)", 10, False, None),
        ("  Daily WET tier = Daily DRY + operator cost/h × 4 × 1.2, rounded to Rs. 50", 10, False, None),
        ("  Per-KM WET tier = DRY + (FULLY WET − DRY) × operator/(fuel+operator), rounded to Rs. 5", 10, False, None),
        ("", 10, False, None),
        ("NOTES", 11, True, None),
        ("  • All rates in LKR, excluding VAT (18%) and SSCL (2.5%).", 10, False, None),
        ("  • 'est.' = no Engineering Cost row for this unit; operator/fuel costs use the "
         "category average (then fleet average). Review before quoting.", 10, False, None),
        ("  • Operator cost 0 (e.g. Motor Bicycles, self-ridden) makes WET equal DRY — intentional.", 10, False, None),
        ("  • Regenerate this file after price updates:  python3 tools/export_rate_matrix.py", 10, False, None),
    ]
    for i, (text, size, bold, color) in enumerate(lines, start=2):
        cell = ws.cell(row=i, column=2, value=text)
        cell.font = Font(name="Calibri", size=size, bold=bold, color=color or INK)
    ws.column_dimensions["B"].width = 110


FLEET_HEADERS = [
    ("#", 5), ("Category", 22), ("E&C No.", 9), ("Model", 24), ("Year", 6),
    ("Registration", 12), ("Fuel /h", 9), ("Operator /h", 10), ("Pricing", 8),
    ("Hourly FULLY WET", 12), ("Hourly WET (driver only)", 12), ("Hourly DRY", 12),
    ("Daily FULLY WET (4h)", 12), ("Daily WET (driver only)", 12), ("Daily DRY (4h)", 12),
    ("Per-KM FULLY WET", 11), ("Per-KM WET (driver only)", 11), ("Per-KM DRY", 11),
]
TIER_FILL = {10: ORANGE, 11: GREEN, 12: GRAY, 13: ORANGE, 14: GREEN, 15: GRAY,
             16: ORANGE, 17: GREEN, 18: GRAY}


def sheet_fleet(wb, fleet):
    ws = wb.create_sheet("Fleet Rate Matrix")
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="FLEET RATE MATRIX 2026 — all 518 units, three rate-basis tiers"
            ).font = Font(bold=True, size=13, color=ORANGE)
    ws.cell(row=2, column=1, value="Rs. excl. VAT · FULLY WET = driver + fuel · WET = driver only, "
            "customer fuels · DRY = vehicle only · Daily = 4-hour block").font = Font(size=9, color=GRAY)

    hdr_row = 4
    for c, (title, width) in enumerate(FLEET_HEADERS, start=1):
        ws.cell(row=hdr_row, column=c, value=title)
        ws.column_dimensions[get_column_letter(c)].width = width
    style_header(ws, hdr_row, len(FLEET_HEADERS), TIER_FILL)

    est_fill = PatternFill("solid", fgColor=ORANGE_SOFT)
    for i, m in enumerate(fleet):
        r = hdr_row + 1 + i
        km = m["r"]["km"] or {}
        values = [
            i + 1, m["cat"], m["ec"] or "", m["model"] or "", m["year"], m["reg"] or "",
            m["fuel"], m["op"], "est." if m.get("est") else "list",
            m["r"]["h"]["fw"], m["r"]["h"]["w"], m["r"]["h"]["d"],
            m["r"]["dy"]["fw"], m["r"]["dy"]["w"], m["r"]["dy"]["d"],
            km.get("fw"), km.get("w"), km.get("d"),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.font = Font(size=10)
            if c >= 7 and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            if c == 9 and m.get("est"):
                cell.fill = est_fill
                cell.font = Font(size=9, italic=True, color="B45309")

    last = hdr_row + len(fleet)
    ws.auto_filter.ref = f"A{hdr_row}:{get_column_letter(len(FLEET_HEADERS))}{last}"
    ws.freeze_panes = f"A{hdr_row + 1}"


def sheet_portable(wb, portable):
    ws = wb.create_sheet("Portable Equipment")
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value="PORTABLE EQUIPMENT DAY-HIRE 2026 — 2-tier"
            ).font = Font(bold=True, size=13, color=ORANGE)
    ws.cell(row=2, column=1, value="Rs./day excl. VAT · WET = incl. fuel & operator · DRY = bare hire"
            ).font = Font(size=9, color=GRAY)

    headers = [("#", 5), ("Category", 26), ("Capacity / Size", 34),
               ("WET Rs./day", 13), ("DRY Rs./day", 13)]
    hdr_row = 4
    for c, (title, width) in enumerate(headers, start=1):
        ws.cell(row=hdr_row, column=c, value=title)
        ws.column_dimensions[get_column_letter(c)].width = width
    style_header(ws, hdr_row, len(headers), {4: GREEN, 5: GRAY})

    for i, p in enumerate(portable):
        r = hdr_row + 1 + i
        for c, v in enumerate([i + 1, p["cat"], p["cap"], p["dw"], p["dd"]], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            cell.font = Font(size=10)
            if c >= 4:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
    last = hdr_row + len(portable)
    ws.auto_filter.ref = f"A{hdr_row}:E{last}"
    ws.freeze_panes = f"A{hdr_row + 1}"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--html", default="index.html")
    ap.add_argument("--out", default="data/Fleet_Rental_Rate_Matrix_2026.xlsx")
    args = ap.parse_args()

    fleet, portable = load_data(args.html)
    wb = Workbook()
    sheet_readme(wb)
    sheet_fleet(wb, fleet)
    sheet_portable(wb, portable)
    wb.save(args.out)
    print(f"WROTE {args.out}  ({len(fleet)} fleet units, {len(portable)} portable items)")


if __name__ == "__main__":
    main()
